import json
import os
from io import BytesIO
from datetime import date, datetime

import requests
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from google.auth.transport.requests import AuthorizedSession
from google.oauth2.service_account import Credentials
from openpyxl import load_workbook


app = FastAPI(title="API Costos y Tarifas ERSeP")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["GET"],
    allow_headers=["*"],
)


FILE_ID = os.getenv("COSTOS_GOOGLE_SHEET_ID")
SERVICE_ACCOUNT_JSON = os.getenv("COSTOS_GOOGLE_SERVICE_ACCOUNT_JSON")


def obtener_credenciales():
    if not FILE_ID:
        raise RuntimeError(
            "Falta la variable COSTOS_GOOGLE_SHEET_ID"
        )

    if not SERVICE_ACCOUNT_JSON:
        raise RuntimeError(
            "Falta la variable COSTOS_GOOGLE_SERVICE_ACCOUNT_JSON"
        )

    credenciales = json.loads(SERVICE_ACCOUNT_JSON)

    return Credentials.from_service_account_info(
        credenciales,
        scopes=[
            "https://www.googleapis.com/auth/drive.readonly"
        ],
    )


def descargar_excel():
    """
    Descarga AJUSTES TARIFARIOS.xlsx directamente desde Google Drive
    y lo abre en memoria. No crea copias ni modifica el archivo original.
    """
    credentials = obtener_credenciales()
    sesion = AuthorizedSession(credentials)

    url = (
        f"https://www.googleapis.com/drive/v3/files/"
        f"{FILE_ID}?alt=media"
    )

    respuesta = sesion.get(url, timeout=30)

    if not respuesta.ok:
        raise RuntimeError(
            f"Google Drive respondió {respuesta.status_code}: "
            f"{respuesta.text}"
        )

    return load_workbook(
        BytesIO(respuesta.content),
        data_only=True,
        read_only=True,
    )


def convertir_valor(valor):
    """
    Convierte valores de Excel a tipos compatibles con JSON.
    """
    if valor is None:
        return ""

    if isinstance(valor, (datetime, date)):
        return valor.isoformat()

    return valor


@app.get("/")
def inicio():
    return {
        "ok": True,
        "servicio": "API Costos y Tarifas ERSeP",
    }


@app.get("/api/hojas")
def listar_hojas():
    """
    Devuelve las pestañas existentes en AJUSTES TARIFARIOS.xlsx.
    """
    try:
        libro = descargar_excel()

        return {
            "ok": True,
            "hojas": libro.sheetnames,
        }

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=str(e),
        )


@app.get("/api/hoja/{nombre_hoja}")
def leer_hoja(nombre_hoja: str):
    """
    Lee una pestaña del Excel.
    La primera fila se utiliza como encabezado.
    """
    try:
        libro = descargar_excel()

        if nombre_hoja not in libro.sheetnames:
            raise HTTPException(
                status_code=404,
                detail=f"No existe la hoja '{nombre_hoja}'",
            )

        hoja = libro[nombre_hoja]

        filas = list(hoja.iter_rows(values_only=True))

        if not filas:
            return {
                "ok": True,
                "hoja": nombre_hoja,
                "cantidad": 0,
                "datos": [],
            }

        encabezados = []

        for i, valor in enumerate(filas[0]):
            if valor is None or str(valor).strip() == "":
                encabezados.append(f"columna_{i + 1}")
            else:
                encabezados.append(str(valor).strip())

        datos = []

        for fila in filas[1:]:

            # Ignorar filas completamente vacías
            if all(valor is None for valor in fila):
                continue

            registro = {}

            for i, encabezado in enumerate(encabezados):
                valor = fila[i] if i < len(fila) else None
                registro[encabezado] = convertir_valor(valor)

            datos.append(registro)

        return {
            "ok": True,
            "hoja": nombre_hoja,
            "cantidad": len(datos),
            "datos": datos,
        }

    except HTTPException:
        raise

    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=str(e),
        )
